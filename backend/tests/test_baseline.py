"""Synthetic, offline baseline. Run from backend with unittest discovery.

Never import main/auth: auth loads local account records at import time.
AI clients are mocked; an outbound socket guard catches accidental calls.
"""

import json
import socket
import unittest
from datetime import date
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock, patch

from docx import Document
from pypdf import PdfWriter

from ai.categorizer import categorize_skills_ai
from evaluator.analyzer import analyze_resume_for_job
from evaluator.scorer import (
    calculate_resume_score, calculate_score_breakdown,
    calculate_total_experience_years, determine_recommendation,
    find_missing_sections,
)
from extractor.extractor import extract_candidate_info
from matcher.matcher import extract_required_years, match_resume_to_job
from parser.document_parser import parse_document
from parser.resume_parser import parse_resume
from utils.resume_checker import looks_like_resume
from utils.skills import find_skills_in_text
from utils.validators import validate_resume_file


def setUpModule():
    guard = patch.object(socket.socket, "connect", side_effect=AssertionError("Network forbidden in tests"))
    guard.start()
    unittest.addModuleCleanup(guard.stop)


def candidate():
    return {
        "name": "Synthetic Candidate", "email": "candidate@example.invalid",
        "phone": "+1 202 555 0100", "skills": ["Python", "SQL"],
        "education": ["Synthetic University, BSc, 2019"],
        "experience": [{"title": "Developer", "duration": "2020-2023",
                        "details": ["Built a synthetic service."]}],
        "detected_sections": ["skills", "education", "experience"],
    }


def completion_client(content):
    client = Mock()
    client.chat.completions.create.return_value = SimpleNamespace(
        choices=[SimpleNamespace(message=SimpleNamespace(content=content))])
    return client


class ExperienceTests(unittest.TestCase):
    def test_absent_or_unparseable_dates(self):
        for value in ({}, {"experience": []}, {"experience": [{"duration": "unknown"}]}):
            with self.subTest(value=value):
                self.assertEqual(calculate_total_experience_years(value), 0)

    def test_single_interval(self):
        self.assertEqual(calculate_total_experience_years(candidate()), 3)

    def test_present_uses_current_year(self):
        with patch("evaluator.scorer.date") as clock:
            clock.today.return_value = date(2026, 10, 19)
            self.assertEqual(calculate_total_experience_years(
                {"experience": [{"duration": "2020-present"}]}), 6)

    def test_overlapping_intervals_not_double_counted(self):
        self.assertEqual(calculate_total_experience_years({"experience": [
            {"duration": "2020-2023"}, {"duration": "2022-2025"}]}), 5)


class CompletenessTests(unittest.TestCase):
    def test_empty_resume_scores_zero(self):
        self.assertEqual(calculate_resume_score({}), 0)
        self.assertEqual(set(find_missing_sections({})),
                         {"name", "email", "phone", "skills", "education", "experience"})

    def test_contact_fields_contribute_individually(self):
        for field, points in (("name", 4), ("email", 3), ("phone", 3)):
            with self.subTest(field=field):
                self.assertEqual(calculate_resume_score({field: "synthetic"}), points)

    def test_complete_example_and_breakdown(self):
        self.assertEqual(calculate_resume_score(candidate()), 67)
        self.assertEqual(find_missing_sections(candidate()), [])
        self.assertEqual(calculate_score_breakdown(candidate()),
                         {"contact": 100, "skills": 40, "education": 100, "experience": 60})

    def test_volume_does_not_exceed_score_boundary(self):
        value = candidate()
        value["skills"] = [f"Synthetic skill {i}" for i in range(100)]
        value["experience"][0]["details"] = ["Synthetic detail"] * 100
        self.assertEqual(calculate_resume_score(value), 100)
        for score in calculate_score_breakdown(value).values():
            self.assertGreaterEqual(score, 0)
            self.assertLessEqual(score, 100)

    def test_recommendation_threshold_edges(self):
        for score, label in ((0, "Not Recommended"), (49, "Not Recommended"),
                             (50, "Consider"), (69, "Consider"), (70, "Recommended"),
                             (84, "Recommended"), (85, "Highly Recommended"),
                             (100, "Highly Recommended")):
            with self.subTest(score=score):
                self.assertEqual(determine_recommendation(score), label)


class MatchingTests(unittest.TestCase):
    def test_case_insensitive_matching_and_missing_skill(self):
        result = match_resume_to_job(candidate(), "PYTHON and Java required")
        self.assertEqual([s.lower() for s in result["matched_skills"]], ["python"])
        self.assertEqual(result["missing_skills"], ["java"])
        self.assertEqual(result["skills_score"], 50)
        self.assertEqual(result["overall_score"], 70)

    def test_word_boundaries(self):
        self.assertEqual(find_skills_in_text("JavaScript and Google", ["java", "go"]), [])

    def test_supported_experience_formats(self):
        for text, expected in (("3+ years", 3), ("3 to 5 years", 3), ("1 year", 1), ("", None)):
            with self.subTest(text=text):
                self.assertEqual(extract_required_years(text), expected)

    def test_matching_score_extremes(self):
        self.assertEqual(match_resume_to_job(candidate(), "Python SQL, 3 years, degree")["overall_score"], 100)
        self.assertEqual(match_resume_to_job({}, "Python, 3 years, degree")["overall_score"], 0)

    def test_unknown_requirements_current_policy(self):
        # Characterizes the current policy, not evidence that this is a valid fit score.
        self.assertEqual(match_resume_to_job(candidate(), "Seeking a talented individual")["overall_score"], 35)


class InputTests(unittest.TestCase):
    def test_upload_validation(self):
        for name, size, valid in ((None, 1, False), ("resume", 1, False),
                                  ("resume.exe", 1, False), ("resume.pdf", 0, False),
                                  ("resume.PDF", 10 * 1024**2, True),
                                  ("resume.docx", 1, True),
                                  ("resume.pdf", 10 * 1024**2 + 1, False)):
            with self.subTest(name=name, size=size):
                result, error = validate_resume_file(name, size)
                self.assertEqual(result, valid)
                self.assertEqual(error is None, valid)

    def test_unsupported_document_types(self):
        for parser in (parse_resume, parse_document):
            with self.subTest(parser=parser.__name__):
                with self.assertRaisesRegex(ValueError, "Unsupported file type"):
                    parser("synthetic.exe", b"synthetic")

    def test_empty_resume_text_rejected(self):
        for text in ("", "   \n"):
            with self.subTest(text=text):
                self.assertFalse(looks_like_resume(text)[0])
                self.assertEqual(extract_candidate_info(text)["experience"], [])

    def test_plain_text_job_description(self):
        self.assertEqual(parse_document("job.txt", b"Python developer"), "Python developer")

    def test_synthetic_docx_paragraph_and_table(self):
        document = Document()
        document.add_paragraph("Synthetic Candidate")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Python"
        stream = BytesIO()
        document.save(stream)
        text = parse_resume("synthetic.docx", stream.getvalue())
        self.assertIn("Synthetic Candidate", text)
        self.assertIn("Python", text)

    def test_blank_pdf_has_no_readable_resume(self):
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        stream = BytesIO()
        writer.write(stream)
        self.assertFalse(looks_like_resume(parse_resume("blank.pdf", stream.getvalue()))[0])

    def test_single_section_extraction(self):
        text = "Synthetic Candidate\ncandidate@example.invalid\nSkills\nPython, SQL\nEducation\nSynthetic University 2019\nExperience\nDeveloper\n2020-2023\nBuilt a service."
        result = extract_candidate_info(text)
        self.assertEqual(result["skills"], ["Python", "SQL"])
        self.assertEqual(len(result["experience"]), 1)
        self.assertEqual(result["experience"][0]["duration"], "2020-2023")
        self.assertTrue(looks_like_resume(text)[0])


class AIFallbackTests(unittest.TestCase):
    def test_resume_only_never_calls_ai(self):
        with patch("evaluator.analyzer.match_resume_to_job_ai") as ai:
            result = analyze_resume_for_job(candidate(), "  ")
        ai.assert_not_called()
        self.assertIsNone(result["job_fit_score"])
        self.assertEqual(result["overall_score"], 67)

    def test_missing_key_and_timeout_fall_back(self):
        expected = match_resume_to_job(candidate(), "Python")
        for failure in (KeyError("OPENROUTER_API_KEY"), TimeoutError("synthetic timeout")):
            with self.subTest(failure=type(failure).__name__):
                with patch("ai.matcher.get_client", side_effect=failure):
                    result = analyze_resume_for_job(candidate(), "Python")
                self.assertEqual(result["job_fit_method"], "rule-based")
                self.assertIsNone(result["job_fit_model"])
                self.assertEqual(result["job_fit_score"], expected["overall_score"])

    def test_invalid_json_and_schema_fall_back(self):
        for raw in ("not JSON", "{}"):
            with self.subTest(raw=raw):
                with patch("ai.matcher.get_client", return_value=completion_client(raw)):
                    result = analyze_resume_for_job(candidate(), "Python")
                self.assertEqual(result["job_fit_method"], "rule-based")
                self.assertEqual(result["job_fit_score"], 100)

    def test_valid_ai_response_is_used(self):
        payload = dict(requirements={"required_skills": ["Python"]}, overall_score=80,
                       match_label="Strong Match", skills_score=80, experience_score=80,
                       education_score=80, experience_note="Synthetic evidence",
                       education_note="Synthetic evidence", matched_skills=["Python"],
                       strengths=["Synthetic strength"], suggestions=["Synthetic follow-up"])
        with patch("ai.matcher.get_client", return_value=completion_client(json.dumps(payload))):
            result = analyze_resume_for_job(candidate(), "Python")
        self.assertEqual(result["job_fit_method"], "ai")
        self.assertEqual(result["overall_score"], 76)
        self.assertEqual(result["strengths"], payload["strengths"])

    def test_weak_match_cap_current_policy(self):
        value = candidate()
        value["skills"] = ["Python", "SQL", "Docker", "Git", "React"]
        value["experience"][0]["details"] *= 10
        with patch("ai.matcher.get_client", side_effect=TimeoutError):
            result = analyze_resume_for_job(value, "Java")
        self.assertEqual(result["resume_quality_score"], 100)
        self.assertEqual(result["job_fit_score"], 40)
        self.assertEqual(result["overall_score"], 49)

    def test_categorizer_rejects_invented_skills(self):
        raw = json.dumps({"groups": [{"category": "Languages", "skills": ["Java"]}]})
        with patch("ai.categorizer.get_client", return_value=completion_client(raw)):
            with self.assertRaisesRegex(ValueError, "preserve"):
                categorize_skills_ai(["Python"])

    def test_empty_skills_skip_provider(self):
        with patch("ai.categorizer.get_client") as client:
            self.assertEqual(categorize_skills_ai([]), {})
        client.assert_not_called()
