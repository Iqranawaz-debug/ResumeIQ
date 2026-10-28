"""Synthetic employment regressions; no files, accounts or providers used."""
import builtins
import importlib
import unittest
from datetime import date
from io import BytesIO, StringIO
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi import UploadFile
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from evaluator.scorer import calculate_total_experience_years, calculate_resume_score
from extractor.extractor import extract_candidate_info
from parser.resume_parser import parse_resume
from utils.employment_dates import interval


def extract(body):
    return extract_candidate_info('Synthetic Candidate\ncandidate@example.invalid\n' + body)


class EmploymentTests(unittest.TestCase):
    def test_single_role(self):
        job = extract('Experience\nEngineer\nExample Ltd.\n2020-2023\nBuilt a service.')['experience'][0]
        self.assertEqual((job['title'], job['organization'], job['duration']), ('Engineer', 'Example Ltd.', '2020-2023'))
        self.assertEqual(job['details'], ['Built a service.'])

    def test_teaching_assistant_regression(self):
        jobs = extract('Experience\nDevelopment Intern\nExample Services\nJuly 2026 – August 2026\nBuilt an agent.\nTeaching Assistant (Computing)\nExample University\nFall 2025\nMentored students; supported grading.\nIT Intern\nExample Authority\nJun 2025 - Aug 2025\nSupported operations.')['experience']
        self.assertEqual([j['title'] for j in jobs], ['Development Intern', 'Teaching Assistant (Computing)', 'IT Intern'])
        self.assertEqual(jobs[1]['details'], ['Mentored students; supported grading.'])
        self.assertEqual(jobs[2]['organization'], 'Example Authority')

    def test_header_orders(self):
        for body in ('Engineer\nExample Ltd.\n2020-2021', 'Example Ltd.\nEngineer\n2020-2021', 'Engineer, Example Ltd. (2020-2021)', 'Example Ltd. | Engineer | 2020-2021'):
            with self.subTest(body=body):
                jobs = extract('Experience\n' + body)['experience']
                self.assertEqual(len(jobs), 1)
                self.assertEqual(jobs[0]['title'], 'Engineer')
                self.assertEqual(jobs[0]['organization'], 'Example Ltd.')

    def test_same_employer_distinct_roles(self):
        jobs = extract('Experience\nEngineer\nExample Ltd.\n2020-2021\nBuilt tools.\nSenior Engineer\nExample Ltd.\n2021-2023\nLed delivery.')['experience']
        self.assertEqual(len(jobs), 2)
        self.assertEqual(jobs[1]['title'], 'Senior Engineer')

    def test_diverse_roles_without_dates(self):
        for title in ('Part-time Nurse', 'Freelance Designer', 'Contract Consultant', 'Teaching Assistant', 'Legal Intern', 'Chef', 'Accountant'):
            with self.subTest(title=title):
                job = extract(f'Experience\n{title}\nExample Ltd.\n• Supported a team.')['experience'][0]
                self.assertEqual(job['title'], title)
                self.assertIsNone(job['duration'])

    def test_missing_organization(self):
        job = extract('Experience\nEngineer\n2020-2021\nBuilt tools.')['experience'][0]
        self.assertIsNone(job['organization'])

    def test_wrapped_title(self):
        job = extract('Experience\nSenior Software\nEngineer\nExample Ltd.\n2020-2021')['experience'][0]
        self.assertEqual(job['title'], 'Senior Software Engineer')

    def test_missing_title(self):
        job = extract('Experience\nExample Hospital\n2020-2021\nSupported operations.')['experience'][0]
        self.assertIsNone(job['title'])
        self.assertEqual(job['organization'], 'Example Hospital')

    def test_ambiguous_employer_preserved(self):
        job = extract('Experience\nEngineer\nExample Collective\n2020-2021')['experience'][0]
        self.assertEqual(job['title'], 'Engineer')
        self.assertEqual(job['header_lines'], ['Example Collective'])
        self.assertIsNone(job['organization'])

    def test_bullet_and_wrap_preserved(self):
        job = extract('Experience\nEngineer\n2020-2021\n• Built a tool for\n  internal users\n• Managed a team\n• Worked with an engineer in 2020')['experience'][0]
        self.assertEqual(job['details'], ['Built a tool for internal users', 'Managed a team', 'Worked with an engineer in 2020'])

    def test_prose_is_not_promoted(self):
        jobs = extract('Experience\nTeacher\n2020-2021\nMentored students\nIT Intern\n2022-2023\nSupported users.')['experience']
        self.assertEqual([j['title'] for j in jobs], ['Teacher', 'IT Intern'])
        self.assertEqual(jobs[0]['details'], ['Mentored students'])

    def test_duplicate_section_records_removed(self):
        self.assertEqual(len(extract(('Experience\nEngineer\n2020-2021\nBuilt tools.\n') * 2)['experience']), 1)


class DateTests(unittest.TestCase):
    def test_written_formats(self):
        for value in ('Jul 2026 - Aug 2026', 'July 2026 – August 2026', 'Jun 2025 — Aug 2025', 'Fall 2025', 'Spring 2025', 'Summer 2026', '2024 - 2025', 'Jan 2024 - Present', '2023 – Current', '06/2025 - 08/2025', 'Jan 2024 − Feb 2024'):
            with self.subTest(value=value):
                self.assertEqual(extract(f'Experience\nEngineer\n{value}')['experience'][0]['duration'], value)

    def test_seasons_do_not_invent_duration(self):
        candidate = extract('Experience\nTeacher\nFall 2025')
        self.assertEqual(candidate['experience'][0]['date_status'], 'uncertain')
        self.assertEqual(calculate_total_experience_years(candidate), 0)

    def test_invalid_dates(self):
        for duration in ('Aug 2025 - Jun 2025', '2026-2025', '13/2025 - 14/2025'):
            with self.subTest(duration=duration):
                self.assertEqual(interval(duration)[1], 'invalid')
                self.assertEqual(calculate_total_experience_years({'experience': [{'duration': duration}]}), 0)

    def test_month_overlap_and_gaps(self):
        candidate = {'experience': [{'duration': s} for s in ('Jan 2020 - Jul 2020', 'Apr 2020 - Oct 2020', 'Jan 2022 - Apr 2022')]}
        self.assertEqual(calculate_total_experience_years(candidate), 1)

    def test_current(self):
        self.assertEqual(interval('Jan 2024 - Current', date(2025, 1, 10))[0], (2024 * 12, 2025 * 12))

    def test_wrapped_date_range(self):
        for value in ('Jan 2024 -\nFeb 2024', 'Jan 2024\n–\nFeb 2024'):
            job = extract(f'Experience\nEngineer\n{value}')['experience'][0]
            self.assertEqual(interval(job['duration'])[0], (2024 * 12, 2024 * 12 + 1))

    def test_numeric_dates_are_not_phone_numbers(self):
        self.assertIsNone(extract('Experience\nEngineer\n06/2025 - 08/2025')['phone'])

    def test_mixed_precision_and_single_date(self):
        for duration in ('2020 - Jan 2021', '2020', 'June 2020'):
            self.assertEqual(interval(duration), (None, 'uncertain'))


class SectionTests(unittest.TestCase):
    def test_alternative_experience_headings(self):
        for heading in ('Work Experience', 'Professional Experience', 'Employment History', 'Career History', 'Internships', 'Relevant Experience', 'PROFESSIONAL EXPERIENCE:', 'Professional\nExperience:'):
            with self.subTest(heading=heading):
                self.assertEqual(len(extract(f'{heading}\nEngineer\n2020-2021')['experience']), 1)

    def test_project_headings(self):
        for heading in ('Projects', 'Academic Projects', 'Personal Projects', 'Featured Projects', 'Portfolio Projects', 'Selected Projects', 'Earlier Projects & Open Source', 'Open Source Contributions', 'Research Projects', 'Open Source', 'EARLIER PROJECTS &\nOPEN SOURCE:'):
            with self.subTest(heading=heading):
                candidate = extract(f'Experience\nEngineer\n2020-2021\nBuilt tools.\n{heading}\nAnalysis Dashboard\nBuilt charts with Python.')
                self.assertEqual(len(candidate['experience']), 1)
                self.assertEqual(candidate['experience'][0]['details'], ['Built tools.'])
                self.assertEqual(candidate['projects'], [{'title': 'Analysis Dashboard', 'details': ['Built charts with Python.']}])
                self.assertIn('Python', candidate['skills'])

    def test_other_sections_bound_experience(self):
        for heading in ('Education', 'Skills', 'Technical Skills', 'Certifications', 'Achievements', 'Publications'):
            with self.subTest(heading=heading):
                candidate = extract(f'Experience\nEngineer\n2020-2021\nBuilt tools.\n{heading}\nExample University 2019')
                self.assertEqual(candidate['experience'][0]['details'], ['Built tools.'])

    def test_empty_sections(self):
        candidate = extract('Experience\nProjects\nEducation\nSkills')
        self.assertEqual(candidate['experience'], [])
        self.assertEqual(candidate['projects'], [])

    def test_ordinary_sentence_is_not_heading(self):
        candidate = extract('Experience\nEngineer\n2020-2021\nProjects supported customer needs.')
        self.assertEqual(candidate['experience'][0]['details'], ['Projects supported customer needs.'])

    def test_docx_table_order(self):
        document = Document()
        document.add_paragraph('Experience')
        document.add_table(rows=1, cols=1).cell(0, 0).text = 'Engineer\n2020-2021\nBuilt tools.'
        document.add_paragraph('Projects')
        document.add_paragraph('Analysis Dashboard')
        document.add_paragraph('Built charts.')
        stream = BytesIO()
        document.save(stream)
        candidate = extract_candidate_info(parse_resume('synthetic.docx', stream.getvalue()))
        self.assertEqual(candidate['experience'][0]['title'], 'Engineer')
        self.assertEqual(candidate['projects'][0]['title'], 'Analysis Dashboard')

    def test_synthetic_pdf_sections(self):
        writer = PdfWriter()
        page = writer.add_blank_page(width=600, height=800)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        lines = ['Experience', 'Engineer', '2020-2021', 'Built tools.', 'Earlier Projects &', 'Open Source', 'Dashboard', 'Built charts.']
        stream.set_data(('BT /F1 12 Tf 20 760 Td ' + ' '.join(f'({line}) Tj 0 -20 Td' for line in lines) + ' ET').encode('ascii'))
        page[NameObject('/Contents')] = writer._add_object(stream)
        output = BytesIO()
        writer.write(output)
        candidate = extract_candidate_info(parse_resume('synthetic.pdf', output.getvalue()))
        self.assertEqual(candidate['experience'][0]['details'], ['Built tools.'])
        self.assertEqual(candidate['projects'][0]['title'], 'Dashboard')

    def test_project_only_has_no_employment_points(self):
        candidate = extract_candidate_info('Projects\nDashboard\nBuilt charts.')
        self.assertEqual(candidate['experience'], [])
        self.assertEqual(calculate_resume_score({**candidate, 'name': None}), 0)


class CompatibilityTests(unittest.TestCase):
    def test_upload_analysis_history_contract(self):
        original = builtins.open
        account_path = Path(__file__).resolve().parents[1] / 'data' / 'users.json'

        def isolated(file, mode='r', *args, **kwargs):
            if isinstance(file, (str, Path)) and Path(file).resolve() == account_path:
                if mode != 'r':
                    raise AssertionError('Account writes forbidden')
                return StringIO('{}')
            return original(file, mode, *args, **kwargs)

        with patch('builtins.open', side_effect=isolated):
            api = importlib.import_module('main')

        async def exercise():
            document = Document()
            for line in ('Synthetic Candidate', 'candidate@example.invalid', 'Skills', 'Python', 'Education', 'Example University 2019', 'Experience', 'Engineer', '2020-2021', 'Built tools.', 'Projects', 'Dashboard', 'Built charts.'):
                document.add_paragraph(line)
            stream = BytesIO()
            document.save(stream)
            memory_file = BytesIO(stream.getvalue())
            memory_file._rolled = False
            result = await api.upload_resume(UploadFile(filename='synthetic.docx', file=memory_file), 'synthetic@example.invalid')
            job = result['candidate']['experience'][0]
            self.assertTrue({'title', 'duration', 'details'} <= set(job))
            self.assertEqual(result['candidate']['projects'][0]['title'], 'Dashboard')
            self.assertEqual(result['extraction_method'], 'rule-based')
            analysis = await api.analyze_resume('', result['resume_id'], None, 'synthetic@example.invalid', use_ai=False)
            self.assertIn('total_experience_years', analysis['analysis'])
            self.assertEqual(len((await api.get_history('synthetic@example.invalid'))['history']), 1)

        async def inline(function, *args, **kwargs):
            return function(*args, **kwargs)

        # Test endpoint data flow without transport/event-loop socket creation.
        # All awaited IO is in-memory and the threadpool boundary runs inline.
        with patch.dict(api.validated_resumes, {}, clear=True), patch.dict(api.upload_versions, {}, clear=True), patch.dict(api.evaluation_histories, {}, clear=True), patch.object(api, 'run_in_threadpool', inline), patch('ai.matcher.get_client', side_effect=AssertionError('No provider calls')):
            coroutine = exercise()
            try:
                with self.assertRaises(StopIteration):
                    coroutine.send(None)
            finally:
                coroutine.close()
