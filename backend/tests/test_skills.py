"""Synthetic regressions, discoverable with both pytest and unittest."""
import unittest
from io import BytesIO
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from extractor.extractor import extract_candidate_info, extract_skills
from parser.resume_parser import parse_resume
from utils.skills import extract_known_skills, categorize_skills
from matcher.matcher import match_resume_to_job


class SkillExtractionTests(unittest.TestCase):
    def test_pdf_wrapped_text_in_memory(self):
        writer = PdfWriter()
        page = writer.add_blank_page(width=600, height=800)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        content = DecodedStreamObject()
        content.set_data(b'BT /F1 12 Tf 20 TL 50 700 Td (Skills) Tj T* (SQL) Tj T* (Server, Python, FastAPI) Tj ET')
        page[NameObject('/Contents')] = content
        stream = BytesIO(); writer.write(stream)
        self.assertEqual(self.extract(parse_resume('synthetic.pdf', stream.getvalue())), ['SQL Server', 'Python', 'FastAPI'])

    def test_aliases_remain_compatible_with_matching(self):
        value = extract_candidate_info('Skills\nsklearn, NLP, REST APIs')
        result = match_resume_to_job(value, 'sklearn, NLP and REST APIs')
        self.assertEqual(result['missing_skills'], [])
        self.assertEqual(result['skills_score'], 100)

    def test_added_vocabulary(self):
        text = 'Skills\nRandom Forest; Hugging Face; Explainable AI; Predictive Modeling; MLflow; RAG; ChromaDB; LLMs; Data Warehousing; ETL; Data Pipelines; Apache Spark; Hadoop; Watchdog'
        self.assertEqual(len(self.extract(text)), 14)

    def extract(self, text):
        return extract_candidate_info(text)['skills']

    def test_requested_technologies(self):
        text = 'Skills\nPython, FastAPI, Scikit-learn, XGBoost, SQL Server, Streamlit, Git, Playwright'
        self.assertEqual(self.extract(text), ['Python', 'FastAPI', 'Scikit-learn', 'XGBoost', 'SQL Server', 'Streamlit', 'Git', 'Playwright'])

    def test_sentence_regressions(self):
        cases = {
            'with local SQLite user accounts.': ['SQLite'],
            'FEATURED PROJECTS (CURRENT WORK)': [],
            'Automation that renames screenshots from their OCR content': ['OCR'],
            'safe dry-run mode and a pytest suite.': ['Pytest'],
            'Extracts charges, taxes and summaries': [],
            'with verification notes': [],
            'Upgraded to a portfolio-grade version': [],
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertEqual(self.extract('Skills\n' + text), expected)
                self.assertNotIn(text, self.extract(text))

    def test_aliases_and_case_deduplication(self):
        self.assertEqual(self.extract('Skills\nsklearn, Scikit-learn, SKLEARN; NLP, Natural Language Processing; REST APIs, REST API'),
                         ['Scikit-learn', 'Natural Language Processing', 'REST API'])

    def test_punctuated_compounds(self):
        self.assertEqual(self.extract('Skills\nC++, C#, .NET; Node.js | CI/CD • IDS/IPS'),
                         ['C++', 'C#', '.NET', 'Node.js', 'CI/CD', 'IDS/IPS'])

    def test_wrapped_phrases_and_soft_hyphens(self):
        self.assertEqual(self.extract('Skills\nScikit-\nlearn, SQL\nServer; Natural Language\nProcessing; Py\u00adthon'),
                         ['Scikit-learn', 'SQL Server', 'Natural Language Processing', 'Python'])

    def test_separate_bullets_do_not_form_compound_skill(self):
        self.assertNotIn('SQL Server', self.extract('Skills\nSQL\n• Server maintenance'))

    def test_lists_do_not_need_commas(self):
        self.assertEqual(self.extract('Skills\nPython\nFastAPI\n• SQLite\n- Pytest'),
                         ['Python', 'FastAPI', 'SQLite', 'Pytest'])

    def test_skill_section_precedes_project_evidence(self):
        self.assertEqual(self.extract('Projects\nBuilt using SQLite.\nSkills\nPython\nProjects\nTested with Playwright.'),
                         ['Python', 'SQLite', 'Playwright'])

    def test_no_skills_section(self):
        self.assertEqual(self.extract('Experience\nBuilt services with FastAPI and SQLite; tested using pytest.'),
                         ['FastAPI', 'SQLite', 'Pytest'])

    def test_substring_false_positives(self):
        self.assertEqual(extract_known_skills('JavaScripted reaction mysqlish githubber internet model evaluation system project dashboard'), [])

    def test_contact_tokens_are_not_skill_evidence(self):
        self.assertEqual(extract_known_skills('python@example.invalid https://github.com/example'), [])

    def test_ambiguous_prose_is_conservative(self):
        self.assertEqual(extract_known_skills('I excel and react swiftly in spring. A watchdog found rust.'), [])
        self.assertEqual(self.extract('Skills\nr, rust, react'), ['R', 'Rust', 'React'])

    def test_nontechnical_domains(self):
        skills = self.extract('Skills\nPatient care; Clinical documentation; Financial modeling; Legal research; Curriculum development; Talent acquisition; SEO; Lead generation')
        groups = categorize_skills(skills)
        self.assertEqual(set(groups), {'Healthcare', 'Finance', 'Legal', 'Education', 'Human Resources', 'Marketing', 'Sales'})

    def test_empty_and_malformed_skill_input(self):
        for value in (None, 12, {}, [], '', '  '):
            with self.subTest(value=value):
                self.assertEqual(extract_known_skills(value), [])
        self.assertEqual(extract_skills(None, [], []), [])
        self.assertEqual(categorize_skills([None, 12, '', '  ']), {})

    def test_exact_categories(self):
        cases = {
            'Python': 'Programming Languages', 'FastAPI': 'Backend',
            'SQL': 'Databases', 'SQLite': 'Databases', 'SQL Server': 'Databases',
            'Pandas': 'Data Science & AI', 'NumPy': 'Data Science & AI',
            'Scikit-learn': 'Data Science & AI', 'XGBoost': 'Data Science & AI',
            'Random Forest': 'Data Science & AI', 'NLP': 'Data Science & AI',
            'Streamlit': 'Backend', 'Power BI': 'Data Science & AI',
            'Tableau': 'Data Science & AI', 'Apache Spark': 'Data Engineering',
            'Data Warehousing': 'Data Engineering', 'Playwright': 'Automation & Testing',
            'Pytest': 'Automation & Testing', 'Tesseract OCR': 'Automation & Testing',
            'Git': 'Development Tools', 'GitHub': 'Development Tools',
            'GitHub Actions': 'Cloud & DevOps', 'Python Automation': 'Automation & Testing',
        }
        for skill, category in cases.items():
            with self.subTest(skill=skill):
                self.assertEqual(list(categorize_skills([skill])), [category])

    def test_category_deduplication_and_existing_unknown_support(self):
        self.assertEqual(categorize_skills(['sklearn', 'Scikit-learn', 'SKLEARN', 'Unlisted specialty']),
                         {'Data Science & AI': ['Scikit-learn'], 'Other': ['Unlisted specialty']})

    def test_longest_phrase_wins(self):
        self.assertEqual(self.extract('Skills\nSQL Server, Python Automation, Tesseract OCR, GitHub Actions'),
                         ['SQL Server', 'Python Automation', 'Tesseract OCR', 'GitHub Actions'])

    def test_docx_multiline_in_memory(self):
        doc = Document()
        doc.add_paragraph('Skills')
        doc.add_paragraph('Scikit-\nlearn, SQL\nServer')
        doc.add_paragraph('Projects')
        doc.add_paragraph('Automation with OCR and pytest.')
        stream = BytesIO(); doc.save(stream)
        self.assertEqual(self.extract(parse_resume('synthetic.docx', stream.getvalue())),
                         ['Scikit-learn', 'SQL Server', 'OCR', 'Pytest'])

    def test_other_candidate_fields_remain_compatible(self):
        result = extract_candidate_info('Synthetic Candidate\ncandidate@example.invalid\nSkills\nPython\nEducation\nSynthetic University 2020\nExperience\nDeveloper\n2020-2023\nBuilt a service.')
        self.assertEqual(result['name'], 'Synthetic Candidate')
        self.assertEqual(result['email'], 'candidate@example.invalid')
        self.assertEqual(result['education'], ['Synthetic University 2020'])
        self.assertEqual(result['experience'][0]['duration'], '2020-2023')
        self.assertEqual(set(result), {'name', 'email', 'phone', 'skills', 'education', 'experience', 'projects', 'detected_sections', 'skill_evidence'})
