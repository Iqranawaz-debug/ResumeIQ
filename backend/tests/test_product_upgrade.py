"""Synthetic tests for the product upgrade: privacy controls (E3), document checks
(E1), requirement-line matching (A1), score impacts (A2), AI reliability (E2) and
grounded AI rewrites (A3). No provider is contacted; AI responses are mocked."""
import builtins
import importlib
import json
from io import BytesIO, StringIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from docx import Document
from fastapi import HTTPException
from pypdf import PdfWriter

from evaluator.analyzer import analyze_resume_for_job, blend_scores
from evaluator.document_checks import document_checks
from evaluator.improvements import bullet_issues, generate_improvements, grounded_rewrite
from evaluator.scorer import calculate_resume_score
from extractor.extractor import extract_candidate_info
from matcher.matcher import match_resume_to_job
from matcher.requirement_lines import evaluate_requirement_lines
from network_guard import OfflineAsyncTestCase, OfflineTestCase, network_guard
from parser.resume_parser import parse_resume
from test_baseline import completion_client
from utils import limits
from utils.candidate_processing import process_resume

RESUME = """Jordan Example
jordan@example.invalid | +1 555 010 0000
Skills
Python, Django, Azure, SQL, Docker
Experience
Software Engineer, Example Corp
Jan 2019 - Present
Worked on the payment API in Python and was responsible for fixing bugs reported by users.
Maintained CI pipelines and Docker deployments on Azure.
Education
BSc Computer Science, Example University, 2018
"""
JD = """About us: We have been building great products for 20 years.
Requirements:
- 5+ years of experience with Python
- Experience with AWS or similar cloud platforms
- Bachelor's degree in Computer Science or related field
- AWS certification
Responsibilities:
- Design and build REST APIs for payment processing
- Mentor junior engineers
Nice to have:
- Experience with Go
We offer great benefits and remote work."""


def resume():
    return extract_candidate_info(RESUME)


def setUpModule():
    global api, recruiter
    original = builtins.open
    account_path = Path(__file__).resolve().parents[1] / 'data' / 'users.json'
    def isolated(file, mode='r', *args, **kwargs):
        if isinstance(file, (str, Path)) and Path(file).resolve() == account_path:
            if mode != 'r':
                raise AssertionError('Account writes forbidden')
            return StringIO('{}')
        return original(file, mode, *args, **kwargs)
    with network_guard(), patch('builtins.open', side_effect=isolated):
        api = importlib.import_module('main')
        recruiter = importlib.import_module('recruiter')


def docx_bytes(build):
    document = Document()
    build(document)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


# E3 ------------------------------------------------------------------------
class PrivacyControlTests(OfflineAsyncTestCase):
    def setUp(self):
        super().setUp()
        for store in (api.validated_resumes, api.evaluation_histories, api.upload_versions, api.last_activity, recruiter.workspaces):
            self.enterContext(patch.dict(store, {}, clear=True))
        self.enterContext(patch.object(recruiter, 'active_users', set()))
        api.validated_resumes['one@example.invalid'] = {'resume_id': 'r', 'candidate': resume(), 'analysis': {'analysis_id': 'a'}}
        api.validated_resumes['two@example.invalid'] = {'resume_id': 'r2', 'candidate': resume()}
        api.evaluation_histories['one@example.invalid'] = [{'name': 'Jordan'}]
        recruiter.workspaces['one@example.invalid'] = {'g': {'candidates': {'c': {}}}}

    async def test_summary_reports_what_is_stored_without_contents(self):
        summary = await api.get_my_data('one@example.invalid')
        self.assertEqual(summary | {'ai_provider_configured': None}, {
            'resume_stored': True, 'analysis_stored': True, 'history_entries': 1, 'versions': 0, 'workspaces': 1,
            'candidates': 1, 'retention_seconds': api.DATA_RETENTION_SECONDS, 'storage': 'memory', 'ai_provider_configured': None})
        self.assertNotIn('Jordan', json.dumps(summary))

    async def test_delete_erases_only_the_callers_data(self):
        await api.delete_my_data('one@example.invalid')
        self.assertNotIn('one@example.invalid', api.validated_resumes)
        self.assertNotIn('one@example.invalid', api.evaluation_histories)
        self.assertNotIn('one@example.invalid', recruiter.workspaces)
        self.assertIn('two@example.invalid', api.validated_resumes)

    async def test_delete_refuses_during_processing(self):
        recruiter.active_users.add('one@example.invalid')
        with self.assertRaises(HTTPException) as error:
            await api.delete_my_data('one@example.invalid')
        self.assertEqual(error.exception.status_code, 409)
        self.assertIn('one@example.invalid', api.validated_resumes)

    async def test_inactive_data_expires_and_active_data_stays(self):
        api.last_activity.update({'one@example.invalid': 0.0, 'two@example.invalid': 10_000.0})
        api.expire_inactive_data(now=api.DATA_RETENTION_SECONDS + 1.0)
        self.assertNotIn('one@example.invalid', api.validated_resumes)
        self.assertNotIn('one@example.invalid', recruiter.workspaces)
        self.assertIn('two@example.invalid', api.validated_resumes)

    async def test_authenticated_requests_refresh_activity(self):
        with patch('main.get_email_from_token', return_value='one@example.invalid'):
            self.assertEqual(api.require_user('Bearer synthetic'), 'one@example.invalid')
        self.assertIn('one@example.invalid', api.last_activity)

    async def test_delete_requires_authentication(self):
        with self.assertRaises(HTTPException) as error:
            api.require_user('')
        self.assertEqual(error.exception.status_code, 401)


# E1 ------------------------------------------------------------------------
class DocumentCheckTests(OfflineTestCase):
    def codes(self, stats, text, candidate=None):
        candidate = candidate if candidate is not None else extract_candidate_info(text)
        return {f['id'] for f in document_checks(stats, text, candidate)['findings']}

    def test_clean_resume_passes_core_checks(self):
        result = document_checks({'format': 'pdf', 'pages': 1, 'image_only_pages': 0}, RESUME * 3, resume())
        self.assertEqual([f['id'] for f in result['findings'] if f['severity'] == 'warning'], [])
        self.assertTrue(result['passed'])
        self.assertIn('not guaranteed', result['note'])

    def test_docx_header_contact_text_boxes_and_tables_are_reported(self):
        def build(document):
            document.sections[0].header.paragraphs[0].text = 'Jordan Example | jordan@example.invalid | +1 555 010 0000'
            document.add_paragraph('Experience')
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = 'Engineer'
            table.cell(0, 1).text = '2020 - 2022'
        stats = {}
        text = parse_resume('synthetic.docx', docx_bytes(build), stats=stats)
        self.assertEqual(stats['tables'], 1)
        self.assertIn('jordan@example.invalid', stats['header_footer_text'])
        self.assertNotIn('jordan@example.invalid', text)
        codes = self.codes(stats, text)
        self.assertTrue({'contact_in_header', 'tables', 'headings'} <= codes)

    def test_pdf_page_stats_and_scanned_message(self):
        writer = PdfWriter(); writer.add_blank_page(width=200, height=200)
        stream = BytesIO(); writer.write(stream)
        stats = {}
        self.assertEqual(parse_resume('blank.pdf', stream.getvalue(), stats=stats), '')
        self.assertEqual((stats['format'], stats['pages'], stats['page_chars']), ('pdf', 1, [0]))
        with self.assertRaises(HTTPException) as error:
            process_resume('scan.pdf', b'x', parser=lambda name, content, stats: stats.update(image_only_pages=1) or '')
        self.assertEqual(error.exception.status_code, 400)
        self.assertIn('scanned', error.exception.detail)

    def test_icon_glyphs_fragmentation_dates_and_missing_contact(self):
        text = 'Experience\nEngineer\nSometime\n' + '  \n' + '\n'.join(['Word'] * 40)
        codes = self.codes({'format': 'docx'}, text)
        self.assertTrue({'icon_glyphs', 'fragmented', 'missing_contact', 'dates'} <= codes)

    def test_user_safe_error_messages(self):
        writer = PdfWriter(); writer.add_blank_page(width=100, height=100); writer.encrypt('synthetic')
        stream = BytesIO(); writer.write(stream)
        for name, content, expected in (('locked.pdf', stream.getvalue(), 'password-protected'),
                                         ('fake.pdf', b'not a pdf', 'not a valid PDF'),
                                         ('fake.docx', b'not a zip', 'not a valid DOCX')):
            with self.subTest(name=name), self.assertRaises(HTTPException) as error:
                process_resume(name, content)
            self.assertEqual(error.exception.status_code, 422)
            self.assertIn(expected, error.exception.detail)

    def test_upload_pipeline_attaches_checks_without_changing_extraction(self):
        content = docx_bytes(lambda d: [d.add_paragraph(line) for line in RESUME.splitlines()])
        candidate = process_resume('synthetic.docx', content)
        self.assertEqual(candidate['document_checks']['format'], 'docx')
        expected = extract_candidate_info(parse_resume('synthetic.docx', content))
        self.assertEqual({k: candidate[k] for k in expected}, expected)


# A1 ------------------------------------------------------------------------
class RequirementLineTests(OfflineTestCase):
    def rows(self):
        return {row['text']: row for row in evaluate_requirement_lines(resume(), JD)['rows']}

    def test_company_and_benefit_prose_is_skipped(self):
        texts = ' '.join(self.rows())
        self.assertNotIn('20 years', texts)
        self.assertNotIn('benefits', texts)

    def test_statuses_reasons_and_basis(self):
        rows = self.rows()
        self.assertEqual(rows['5+ years of experience with Python']['status'], 'matched')
        self.assertIn('not calculated', rows['5+ years of experience with Python']['reason'])
        aws = rows['Experience with AWS or similar cloud platforms']
        self.assertEqual((aws['status'], aws['basis']), ('partial', 'inference'))
        self.assertIn('AWS', aws['reason'])
        degree = rows["Bachelor's degree in Computer Science or related field"]
        self.assertEqual((degree['type'], degree['status']), ('education', 'matched'))
        self.assertIn('not verified', degree['reason'])
        self.assertEqual(rows['AWS certification']['status'], 'not_evidenced')
        self.assertEqual(rows['Mentor junior engineers']['status'], 'unclear')
        self.assertEqual(rows['Experience with Go']['context'], 'preferred')

    def test_wording_overlap_is_labelled_inference_with_a_real_quote(self):
        jd = 'Responsibilities:\n- Maintain CI pipelines and deployments for the platform'
        row = evaluate_requirement_lines(resume(), jd)['rows'][0]
        self.assertEqual((row['status'], row['basis']), ('partial', 'inference'))
        self.assertIn(row['evidence'][0]['excerpt'], RESUME)

    def test_years_below_requirement_and_missing_education(self):
        candidate = resume(); candidate['education'] = []
        rows = evaluate_requirement_lines(candidate, 'Requirements:\n- 12+ years of experience\n- Master degree required')['rows']
        self.assertEqual([r['status'] for r in rows], ['partial', 'not_evidenced'])

    def test_display_only_and_attached_to_explanation(self):
        with patch('evaluator.analyzer.match_resume_to_job_ai', side_effect=AssertionError):
            analysis = analyze_resume_for_job(resume(), JD, use_ai=False)
        self.assertEqual(analysis['job_fit_score'], match_resume_to_job(resume(), JD)['overall_score'])
        self.assertEqual(analysis['explanation']['requirement_lines']['counts']['matched'], 2)
        self.assertIsNone(analyze_resume_for_job(resume(), '', use_ai=False)['explanation']['requirement_lines'])

    def test_bounded(self):
        jd = 'Requirements:\n' + '\n'.join(f'- Experience with tool{i} systems' for i in range(80))
        result = evaluate_requirement_lines(resume(), jd)
        self.assertEqual(len(result['rows']), 40)
        self.assertTrue(result['bounded'])


# A2 ------------------------------------------------------------------------
class ScoreImpactTests(OfflineTestCase):
    def test_estimates_match_rerunning_the_real_scoring(self):
        candidate = resume()
        jd = 'Requirements:\n- Python, Kubernetes, Terraform\nNice to have:\n- Redis'
        analysis = analyze_resume_for_job(candidate, jd, use_ai=False)
        items = {item['action']: item for item in analysis['impacts']['items']}
        k8s = items['Show your Kubernetes experience in the resume']
        edited = dict(candidate, skills=candidate['skills'] + ['Kubernetes'])
        rerun = analyze_resume_for_job(edited, jd, use_ai=False)
        self.assertEqual(k8s['overall_delta'], rerun['overall_score'] - analysis['overall_score'])
        self.assertEqual(k8s['job_fit_delta'], rerun['job_fit_score'] - analysis['job_fit_score'])
        self.assertGreater(k8s['job_fit_delta'], items['Show your Redis experience in the resume']['job_fit_delta'])
        self.assertIn('Only if', k8s['condition'])

    def test_quality_only_without_job_description(self):
        candidate = {'name': 'A', 'skills': ['Python'], 'experience': [{'details': []}]}
        result = analyze_resume_for_job(candidate, '', use_ai=False)
        actions = {item['action']: item for item in result['impacts']['items']}
        self.assertEqual(actions['Add an email address as plain text near the top']['overall_delta'], 3)
        self.assertEqual(actions['List up to 4 more skills you genuinely use']['quality_delta'], 20)
        self.assertEqual(actions['Add your education, including completion status']['quality_delta'], 20)
        self.assertTrue(all(item['job_fit_delta'] == 0 for item in result['impacts']['items']))

    def test_ai_scores_only_estimate_quality_changes(self):
        payload = dict(requirements={'required_skills': ['Python']}, overall_score=80, match_label='Strong Match',
                       skills_score=80, experience_score=80, education_score=80, experience_note='x',
                       education_note='y', matched_skills=['Python'])
        candidate = {'skills': ['Python'], 'experience': []}
        with patch('ai.matcher.get_client', return_value=completion_client(json.dumps(payload))):
            result = analyze_resume_for_job(candidate, 'Python, Kubernetes')
        self.assertTrue(result['impacts']['scope_note'])
        self.assertTrue(all(item['kind'] == 'quality' for item in result['impacts']['items']))
        email = next(i for i in result['impacts']['items'] if 'email' in i['action'])
        self.assertEqual(email['overall_delta'], blend_scores(calculate_resume_score(candidate) + 3, {'overall_score': 80, 'match_label': 'Strong Match'})[1] - result['overall_score'])

    def test_no_actions_for_a_complete_resume(self):
        candidate = {'name': 'A', 'email': 'a@example.invalid', 'phone': '1', 'skills': list('abcde'),
                     'education': ['x'], 'experience': [{'details': ['d'] * 10}]}
        self.assertEqual(analyze_resume_for_job(candidate, '', use_ai=False)['impacts']['items'], [])


# E2 ------------------------------------------------------------------------
class AIReliabilityTests(OfflineTestCase):
    payload = dict(requirements={'required_skills': ['Python', 'Terraform', 'Kubernetes']}, overall_score=72,
                   match_label='Partial Match', skills_score=70, experience_score=90, education_score=100,
                   experience_note='x', education_note='y', matched_skills=['Python', 'Terraform'], missing_skills=['Kubernetes'])

    def client(self, *contents):
        client = Mock()
        client.chat.completions.create.side_effect = [SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=c))]) for c in contents]
        return client

    def test_ai_score_stays_main_and_local_score_is_beside_it(self):
        with patch('ai.matcher.get_client', return_value=self.client(json.dumps(self.payload))):
            result = analyze_resume_for_job(resume(), JD)
        local = match_resume_to_job(resume(), JD)
        self.assertEqual((result['ai_status'], result['job_fit_score']), ('used', 72))
        self.assertEqual(result['local_job_fit_score'], local['overall_score'])
        self.assertEqual(result['local_overall_score'], blend_scores(result['resume_quality_score'], local)[1])
        self.assertEqual(result['explanation']['local_comparison']['job_fit_score'], local['overall_score'])

    def test_ai_matches_are_labelled_explicit_or_inferred(self):
        with patch('ai.matcher.get_client', return_value=self.client(json.dumps(self.payload))):
            result = analyze_resume_for_job(resume(), JD)
        self.assertEqual(result['ai_match_evidence'], {'Python': 'explicit', 'Terraform': 'inferred'})

    def test_one_retry_for_invalid_output_then_fallback(self):
        client = self.client('not json', json.dumps(self.payload))
        with patch('ai.matcher.get_client', return_value=client):
            self.assertEqual(analyze_resume_for_job(resume(), JD)['ai_status'], 'used')
        self.assertEqual(client.chat.completions.create.call_count, 2)
        client = self.client('not json', '{}')
        with patch('ai.matcher.get_client', return_value=client):
            result = analyze_resume_for_job(resume(), JD)
        self.assertEqual((result['ai_status'], result['job_fit_method']), ('fallback', 'rule-based'))
        self.assertIn('unusable', result['processing_note'])

    def test_no_retry_for_timeouts_and_clear_reasons(self):
        client = Mock(); client.chat.completions.create.side_effect = TimeoutError
        with patch('ai.matcher.get_client', return_value=client):
            result = analyze_resume_for_job(resume(), JD)
        self.assertEqual(client.chat.completions.create.call_count, 1)
        self.assertIn('unavailable', result['processing_note'])
        with patch('ai.matcher.get_client', side_effect=KeyError('OPENROUTER_API_KEY')):
            self.assertIn('no AI provider is configured', analyze_resume_for_job(resume(), JD)['processing_note'])

    def test_local_requests_report_status(self):
        self.assertEqual(analyze_resume_for_job(resume(), JD, use_ai=False)['ai_status'], 'not_requested')
        self.assertEqual(analyze_resume_for_job(resume(), '', use_ai=False)['ai_status'], 'not_applicable')


# A3 ------------------------------------------------------------------------
class GroundedRewriteTests(OfflineTestCase):
    ORIGINAL = 'Worked on the payment API in Python and was responsible for fixing bugs reported by users.'

    def test_accepts_honest_rewording_with_placeholder(self):
        for text in ('Developed the payment API in Python and fixed bugs reported by users [result, if known].',
                     'Fixed user-reported bugs and developed the payment API in Python.'):
            with self.subTest(text=text):
                self.assertEqual(grounded_rewrite(self.ORIGINAL, text), (True, None))

    def test_rejects_invented_facts(self):
        cases = {'Led development of the payment API in Python.': 'responsibility',
                 'Developed the payment API in Python and Django.': 'tool',
                 'Developed the payment API in Python, fixing 40 bugs.': 'number',
                 'Developed the payment API at Stripe in Python.': 'name',
                 'Developed the payment API in Python, significantly boosting revenue.': 'new content',
                 'Developed the payment API in Python [saved 20 hours].': 'placeholder',
                 'Developed the payment API in Python; did not fix bugs.': 'negation',
                 self.ORIGINAL: 'unchanged', '': 'empty'}
        for text, reason in cases.items():
            with self.subTest(text=text):
                accepted, message = grounded_rewrite(self.ORIGINAL, text)
                self.assertFalse(accepted)
                self.assertIn(reason, message)

    def test_bullet_issues(self):
        self.assertEqual([i['code'] for i in bullet_issues(self.ORIGINAL)], ['weak_opening', 'no_outcome'])
        self.assertEqual(bullet_issues('Reduced report preparation time by 30% by automating SQL queries.'), [])
        self.assertIn('first_person', [i['code'] for i in bullet_issues('I was assigned to fix the login page.')])

    def test_local_review_lists_flagged_bullets(self):
        review = generate_improvements(resume(), {})['bullet_review']
        self.assertEqual(review[0]['original_text'], self.ORIGINAL)
        self.assertEqual(review[0]['issues'][0]['code'], 'weak_opening')

    def ai(self, rewrites):
        return completion_client(json.dumps({'rewrites': rewrites}))

    def test_suggestions_merge_only_validated_rewrites(self):
        from ai.rewriter import suggest_rewrites
        candidate = resume(); local = generate_improvements(candidate, {})
        good = 'Developed the payment API in Python and fixed bugs reported by users [result, if known].'
        client = self.ai([{'index': 0, 'text': good, 'rationale': 'Direct verb.'},
                          {'index': 1, 'text': 'Led CI pipelines and Docker deployments on Azure for 50 teams.'}])
        with patch('ai.rewriter.get_client', return_value=client):
            result = suggest_rewrites(candidate, local)
        item = next(i for i in result['suggestions'] if i['original_text'] == self.ORIGINAL)
        self.assertEqual((item['suggested_text'], item['source'], item['has_placeholder']), (good, 'ai-grounded', True))
        self.assertEqual((result['ai']['accepted'], len(result['ai']['rejected'])), (1, 1))
        self.assertEqual(result['method'], 'ai-assisted')
        self.assertNotIn('ai-grounded', json.dumps(local))

    def test_payload_privacy_and_failures(self):
        from ai.rewriter import suggest_rewrites
        candidate = resume(); local = generate_improvements(candidate, {})
        client = self.ai([])
        with patch('ai.rewriter.get_client', return_value=client):
            suggest_rewrites(candidate, local)
        sent = json.loads(client.chat.completions.create.call_args.kwargs['messages'][1]['content'])
        self.assertEqual(set(sent), {'bullets'})
        for private in ('Jordan', 'jordan@example.invalid', 'Example Corp', 'Example University'):
            self.assertNotIn(private, json.dumps(sent))
        for failure in (TimeoutError('secret'), KeyError('secret')):
            with self.subTest(failure=failure), patch('ai.rewriter.get_client', side_effect=failure):
                result = suggest_rewrites(candidate, local)
            self.assertEqual(result['method'], 'local')
            self.assertNotIn('secret', json.dumps(result))
        with patch('ai.rewriter.get_client', return_value=completion_client('{"rewrites":[{"index":0,"text":"x","extra":1}]}')):
            self.assertEqual(suggest_rewrites(candidate, local)['ai']['accepted'], 0)

    def test_organization_bullets_are_not_sent(self):
        from ai.rewriter import eligible_bullets
        candidate = resume()
        candidate['experience'][0]['details'][0] = 'Worked on Example Corp billing in Python.'
        local = generate_improvements(candidate, {})
        self.assertNotIn('Example Corp', json.dumps(eligible_bullets(candidate, local)))


class ImprovementEndpointAITests(OfflineAsyncTestCase):
    async def test_ai_mode_is_cached_separately_and_local_is_reused(self):
        record = {'resume_id': 'r', 'candidate': resume(), 'analysis': {'analysis_id': 'a', 'explanation': {}}}
        self.enterContext(patch.dict(api.validated_resumes, {'one@example.invalid': record}, clear=True))
        with patch('main.suggest_rewrites', side_effect=lambda c, local: dict(local, method='ai-assisted')) as ai:
            local = await api.improve_resume('r', 'a', 'one@example.invalid', use_ai=False)
            first = await api.improve_resume('r', 'a', 'one@example.invalid', use_ai=True)
            second = await api.improve_resume('r', 'a', 'one@example.invalid', use_ai=True)
        ai.assert_called_once()
        self.assertEqual(first, second)
        self.assertEqual(local['improvements']['method'], 'local')
        self.assertEqual(first['improvements']['method'], 'ai-assisted')
