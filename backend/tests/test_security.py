"""Phase 2 regression tests. Account storage is isolated before app import."""
import builtins
import importlib
import json
import unittest
from io import BytesIO, StringIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile

from fastapi import HTTPException, UploadFile
from docx import Document
from pypdf import PdfWriter
from pydantic import ValidationError
from ai.schemas import AIMatchResult
from ai.privacy import candidate_payload, redact_text
from evaluator.analyzer import analyze_resume_for_job
from parser.document_parser import parse_document
from utils import limits
from utils.upload import read_upload
from utils.body_limit import BodyLimitMiddleware
from test_baseline import candidate, completion_client
from network_guard import OfflineAsyncTestCase, OfflineTestCase, network_guard


def setUpModule():
    global api, auth
    original = builtins.open
    account_path = Path(__file__).resolve().parents[1] / 'data' / 'users.json'
    def isolated(file, mode='r', *args, **kwargs):
        if isinstance(file, (str, Path)) and Path(file).resolve() == account_path:
            if mode != 'r':
                raise AssertionError('Account writes forbidden')
            return StringIO('{}')
        return original(file, mode, *args, **kwargs)
    with network_guard(), patch('builtins.open', side_effect=isolated):
        auth = importlib.import_module('auth.auth')
        api = importlib.import_module('main')
    for target in ('auth.auth._save_users',):
        guard = patch(target, side_effect=AssertionError('External effects forbidden'))
        guard.start()
        unittest.addModuleCleanup(guard.stop)


def upload(name, data, size=None):
    return UploadFile(filename=name, file=BytesIO(data), size=size)


class SessionTests(OfflineAsyncTestCase):
    def setUp(self):
        super().setUp()
        for obj in (auth.users, auth.sessions, auth.login_attempts):
            context = patch.dict(obj, {}, clear=True)
            context.start()
            self.addCleanup(context.stop)
        auth.users['test@example.invalid'] = {'salt': 'synthetic', 'password_hash': auth._hash_password('password', 'synthetic')}
        self.request = SimpleNamespace(client=SimpleNamespace(host='synthetic-client'))

    async def test_login_normalizes_existing_account_and_logout_revokes(self):
        result = await api.login(api.LoginRequest(email=' TEST@EXAMPLE.INVALID ', password='password'), self.request)
        self.assertEqual(result['username'], 'test@example.invalid')
        self.assertEqual(api.require_user('Bearer ' + result['token']), 'test@example.invalid')
        await api.logout('Bearer ' + result['token'])
        with self.assertRaises(HTTPException) as error:
            api.require_user('Bearer ' + result['token'])
        self.assertEqual(error.exception.status_code, 401)

    async def test_session_expires_at_boundary(self):
        with patch('auth.auth.time.monotonic', return_value=100):
            token = auth.create_session('TEST@example.invalid')
        with patch('auth.auth.time.monotonic', return_value=100 + auth.SESSION_SECONDS - 1):
            self.assertEqual(auth.get_email_from_token(token), 'test@example.invalid')
        with patch('auth.auth.time.monotonic', return_value=100 + auth.SESSION_SECONDS):
            self.assertIsNone(auth.get_email_from_token(token))

    async def test_login_throttle_and_recovery(self):
        with patch('auth.auth.LOGIN_ATTEMPTS', 2), patch('auth.auth.time.monotonic', return_value=100):
            for _ in range(2):
                with self.assertRaises(HTTPException) as error:
                    await api.login(api.LoginRequest(email='test@example.invalid', password='wrong'), self.request)
                self.assertEqual(error.exception.status_code, 401)
            with self.assertRaises(HTTPException) as error:
                await api.login(api.LoginRequest(email='test@example.invalid', password='password'), self.request)
            self.assertEqual(error.exception.status_code, 429)
        with patch('auth.auth.time.monotonic', return_value=100 + auth.LOGIN_WINDOW):
            self.assertTrue(auth.allow_login('synthetic-client'))

    async def test_unknown_and_wrong_password_have_same_message(self):
        messages = []
        for email in ('test@example.invalid', 'unknown@example.invalid'):
            with self.assertRaises(HTTPException) as error:
                await api.login(api.LoginRequest(email=email, password='wrong'), self.request)
            messages.append(error.exception.detail)
        self.assertEqual(messages[0], messages[1])


class UploadTests(OfflineAsyncTestCase):
    def setUp(self):
        super().setUp()
        for state in (auth.users, auth.sessions, auth.login_attempts,
                      api.validated_resumes, api.upload_versions, api.evaluation_histories):
            self.enterContext(patch.dict(state, {}, clear=True))

    async def test_real_http_workflow_and_safe_validation_errors(self):
        async def request(path, body=b'', content_type='application/json', token='', method='POST'):
            sent = []
            async def receive():
                return {'type': 'http.request', 'body': body, 'more_body': False}
            async def send(message):
                sent.append(message)
            scope = {'type': 'http', 'asgi': {'version': '3.0'}, 'http_version': '1.1',
                     'method': method, 'scheme': 'http', 'path': path, 'raw_path': path.encode(),
                     'query_string': b'', 'root_path': '', 'server': ('test', 80),
                     'client': ('synthetic', 1234), 'headers': [(b'content-type', content_type.encode()),
                     (b'authorization', ('Bearer ' + token).encode())]}
            await api.app(scope, receive, send)
            return sent[0]['status'], json.loads(b''.join(m.get('body', b'') for m in sent[1:]))
        with patch.dict(auth.users, {}, clear=True), patch.dict(auth.sessions, {}, clear=True), patch.dict(auth.login_attempts, {}, clear=True), patch.dict(api.validated_resumes, {}, clear=True), patch.dict(api.evaluation_histories, {}, clear=True), patch('auth.auth._save_users'):
            status, result = await request('/api/signup', json.dumps({'email': 'SYNTHETIC@example.invalid', 'password': 'synthetic-password'}).encode())
            self.assertEqual(status, 200)
            token = result['token']
            doc = Document()
            for line in ['Synthetic Candidate', 'test@example.invalid', 'Skills', 'Python, SQL', 'Education', 'Synthetic University 2020']:
                doc.add_paragraph(line)
            stream = BytesIO(); doc.save(stream)
            body = b'--test-boundary\r\nContent-Disposition: form-data; name="file"; filename="synthetic.docx"\r\nContent-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n' + stream.getvalue() + b'\r\n--test-boundary--\r\n'
            with patch('ai.categorizer.categorize_skills_ai', side_effect=TimeoutError):
                status, result = await request('/api/upload', body, 'multipart/form-data; boundary=test-boundary', token)
            self.assertEqual(status, 200)
            status, result = await request('/api/analyze', ('resume_id=' + result['resume_id']).encode(), 'application/x-www-form-urlencoded', token)
            self.assertEqual(status, 200)
            self.assertIsNone(result['analysis']['job_fit_score'])
            status, result = await request('/api/history', token=token, method='GET')
            self.assertEqual(len(result['history']), 1)
            self.assertEqual((await request('/api/logout', token=token))[0], 200)
            self.assertEqual((await request('/api/history', token=token, method='GET'))[0], 401)
            status, result = await request('/api/login', b'{"email": [], "password": "synthetic-private"}')
            self.assertEqual(status, 422)
            self.assertNotIn('synthetic-private', json.dumps(result))

    async def test_extension_empty_and_byte_limits(self):
        for name, content, maximum, expected in [('bad.exe', b'abc', 3, 415), ('ok.txt', b'', 3, 422), ('ok.txt', b'abcd', 3, 413)]:
            with self.subTest(name=name, content=content):
                with self.assertRaises(HTTPException) as error:
                    await read_upload(upload(name, content), ('txt',), maximum)
                self.assertEqual(error.exception.status_code, expected)
        self.assertEqual(await read_upload(upload('ok.TXT', b'abc'), ('txt',), 3), b'abc')

    async def test_known_oversize_not_read(self):
        file = upload('ok.txt', b'abc', size=100)
        with self.assertRaises(HTTPException):
            await read_upload(file, ('txt',), 3)
        self.assertEqual(file.file.tell(), 0)

    async def test_unknown_size_read_stops_at_limit_plus_one(self):
        file = upload('ok.txt', b'x' * 100)
        with self.assertRaises(HTTPException):
            await read_upload(file, ('txt',), 3)
        self.assertEqual(file.file.tell(), 4)

    async def test_invalid_signatures_and_text_encoding(self):
        for name, content in [('bad.pdf', b'not pdf'), ('bad.docx', b'not zip'), ('bad.txt', b'\xff'), ('bad.txt', b'a\x00b')]:
            with self.subTest(name=name, content=content):
                with self.assertRaises(ValueError):
                    parse_document(name, content)

    async def test_zip_entry_and_expanded_limits(self):
        content = BytesIO()
        with ZipFile(content, 'w') as archive:
            archive.writestr('[Content_Types].xml', 'x')
            archive.writestr('word/document.xml', 'x' * 100)
        for limit, value in [('ZIP_ENTRIES', 1), ('ZIP_BYTES', 50)]:
            with self.subTest(limit=limit), patch.object(limits, limit, value):
                with self.assertRaises(limits.DocumentLimitError):
                    parse_document('large.docx', content.getvalue())

    async def test_pdf_pages_and_encryption(self):
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.add_blank_page(width=100, height=100)
        content = BytesIO(); writer.write(content)
        with patch.object(limits, 'PDF_PAGES', 1):
            with self.assertRaises(limits.DocumentLimitError):
                parse_document('large.pdf', content.getvalue())
        writer.encrypt('synthetic')
        content = BytesIO(); writer.write(content)
        with self.assertRaisesRegex(ValueError, 'Encrypted'):
            parse_document('encrypted.pdf', content.getvalue())

    async def test_text_and_docx_structure_limits(self):
        with patch.object(limits, 'TEXT_CHARS', 3):
            with self.assertRaises(limits.DocumentLimitError):
                parse_document('large.txt', b'abcd')
        doc = Document(); doc.add_paragraph('Synthetic')
        content = BytesIO(); doc.save(content)
        with patch.object(limits, 'XML_NODES', 1):
            with self.assertRaises(limits.DocumentLimitError):
                parse_document('large.docx', content.getvalue())

    async def test_body_limit_declared_and_chunked_before_app(self):
        async def downstream(*args):
            self.fail('Oversized request reached application')
        for headers in ([(b'content-length', b'100')], []):
            sent = []
            async def receive():
                return {'type': 'http.request', 'body': b'abcd', 'more_body': False}
            async def send(message):
                sent.append(message)
            with patch.object(limits, 'REQUEST_BYTES', 3):
                await BodyLimitMiddleware(downstream)({'type': 'http', 'headers': headers}, receive, send)
            self.assertEqual(sent[0]['status'], 413)

    async def test_upload_retains_no_raw_document_and_analysis_requires_id(self):
        with patch.dict(api.validated_resumes, {}, clear=True), patch.dict(api.upload_versions, {}, clear=True):
            with patch('main.parse_resume', return_value='synthetic'), patch('main.looks_like_resume', return_value=(True, None)), patch('main.extract_candidate_info', return_value=candidate()), patch('ai.categorizer.categorize_skills_ai', side_effect=TimeoutError):
                result = await api.upload_resume(upload('synthetic.pdf', b'%PDF-synthetic'), 'test@example.invalid')
            stored = api.validated_resumes['test@example.invalid']
            self.assertNotIn('bytes', stored)
            self.assertNotIn('text', stored)
            self.assertEqual(result['categorization_method'], 'rule-based')
            with self.assertRaises(HTTPException) as error:
                await api.analyze_resume('', 'stale-id', None, 'test@example.invalid', use_ai=False)
            self.assertEqual(error.exception.status_code, 409)
            with self.assertRaises(HTTPException) as error:
                await api.analyze_resume('', result['resume_id'], upload('job.txt', b'  '), 'test@example.invalid', use_ai=False)
            self.assertEqual(error.exception.status_code, 422)
            with self.assertRaises(HTTPException) as error:
                await api.analyze_resume('hidden', result['resume_id'], upload('job.txt', b'Python'), 'test@example.invalid', use_ai=False)
            self.assertEqual(error.exception.status_code, 400)


class PrivacyAndAITests(OfflineTestCase):
    def test_payload_removes_contact_fields_and_embedded_identifiers(self):
        value = candidate()
        value['experience'][0]['details'] = [value['name'] + ' ' + value['email'] + ' ' + value['phone']]
        payload = candidate_payload(value)
        for field in ('name', 'email', 'phone'):
            self.assertNotIn(field, payload)
            self.assertNotIn(value[field], json.dumps(payload))
        self.assertNotIn(value['email'], redact_text('Contact ' + value['email'], value))

    def test_all_ai_score_fields_are_strict_and_bounded(self):
        payload = dict(requirements={}, overall_score=80, match_label='Strong Match', skills_score=80, experience_score=80, education_score=80, experience_note='synthetic', education_note='synthetic')
        for field in ('overall_score', 'skills_score', 'experience_score', 'education_score'):
            for invalid in (-1, 101, '80', True, 80.5):
                with self.subTest(field=field, invalid=invalid), self.assertRaises(ValidationError):
                    AIMatchResult(**(payload | {field: invalid}))

    def test_contradictory_labels_and_skill_lists_fall_back(self):
        payload = dict(requirements={'required_skills': ['Python']}, overall_score=80, match_label='Strong Match', skills_score=80, experience_score=80, education_score=80, experience_note='synthetic', education_note='synthetic')
        for change in ({'match_label': 'Weak Match'}, {'matched_skills': ['Python'], 'missing_skills': ['Python']}, {'matched_skills': ['Invented']}):
            with self.subTest(change=change), patch('ai.matcher.get_client', return_value=completion_client(json.dumps(payload | change))):
                result = analyze_resume_for_job(candidate(), 'Python')
            self.assertEqual(result['job_fit_method'], 'rule-based')
            self.assertIn('local', result['processing_note'])
